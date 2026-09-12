"""Regression checks for lossless, narrowly scoped corpus corrections."""
import copy
import io
import unittest
import zipfile

from docx import Document
from lxml import etree

from audit_docx_corpus import NS, sha
from repair_docx_text import apply_text_plan
from repair_docx_jurisdiction_table import apply_plan


class SourcedRepairTests(unittest.TestCase):
    def fixture(self):
        doc = Document()
        doc.add_paragraph('Original statement')
        doc.add_paragraph('Preserve this separate generation detail')
        table = doc.add_table(rows=2, cols=3)
        for cell, value in zip(table.rows[0].cells, ['State', 'Agency', 'Website']):
            cell.text = value
        for cell, value in zip(table.rows[1].cells, ['AL', 'Wrong agency', 'Website']):
            cell.text = value
        doc.sections[0].footer.paragraphs[0].text = 'Footer evidence'
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def text_plan(self, source):
        return {'source_sha256': sha(source), 'edits': [{
            'xpath': '/w:document/w:body/w:p[1]',
            'before': 'Original statement', 'after': 'Corrected statement',
            'source_url': 'https://official.example.gov/evidence'}]}

    def table_plan(self, source):
        with zipfile.ZipFile(io.BytesIO(source)) as z:
            root = etree.fromstring(z.read('word/document.xml'))
        table = root.xpath('/w:document/w:body/w:tbl', namespaces=NS)[0]
        return {'source_sha256': sha(source), 'paragraph_edits': [],
                'table_xpath': '/w:document/w:body/w:tbl',
                'table_xml_sha256': sha(etree.tostring(table)),
                'old_row_count': 1, 'old_headers': ['State', 'Agency', 'Website'],
                'headers': ['State', 'Agency', 'Website'],
                'rows': [{'values': ['AL', 'Sourced agency', 'Agency website'],
                          'source_url': 'https://official.example.gov/evidence',
                          'links': {'2': 'https://official.example.gov/help'}}]}

    def assert_other_parts_unchanged(self, source, result):
        with zipfile.ZipFile(io.BytesIO(source)) as old, zipfile.ZipFile(io.BytesIO(result)) as new:
            self.assertEqual(old.namelist(), new.namelist())
            for name in old.namelist():
                if name not in ('word/document.xml', 'word/_rels/document.xml.rels'):
                    self.assertEqual(old.read(name), new.read(name), name)

    def test_text_repair_preserves_other_paragraphs_and_package_parts(self):
        source = self.fixture()
        result = apply_text_plan(source, self.text_plan(source))
        self.assert_other_parts_unchanged(source, result)
        doc = Document(io.BytesIO(result))
        self.assertEqual([p.text for p in doc.paragraphs],
                         ['Corrected statement', 'Preserve this separate generation detail'])
        self.assertEqual(doc.tables[0].cell(1, 1).text, 'Wrong agency')

    def test_text_repair_rejects_wrong_generation_and_stale_target(self):
        source = self.fixture()
        plan = self.text_plan(source)
        plan['source_sha256'] = '0' * 64
        with self.assertRaisesRegex(ValueError, 'Source hash changed'):
            apply_text_plan(source, plan)
        plan = self.text_plan(source)
        plan['edits'][0]['before'] = 'A different generation'
        with self.assertRaisesRegex(ValueError, 'Target paragraph changed'):
            apply_text_plan(source, plan)

    def test_text_repair_rejects_missing_evidence_and_repeated_location(self):
        source = self.fixture()
        plan = self.text_plan(source)
        plan['edits'][0]['source_url'] = ''
        with self.assertRaisesRegex(ValueError, 'Missing correction evidence'):
            apply_text_plan(source, plan)
        plan = self.text_plan(source)
        plan['edits'].append(copy.deepcopy(plan['edits'][0]))
        with self.assertRaisesRegex(ValueError, 'Repeated correction location'):
            apply_text_plan(source, plan)

    def test_table_repair_preserves_narrative_and_other_parts(self):
        source = self.fixture()
        result = apply_plan(source, self.table_plan(source))
        self.assert_other_parts_unchanged(source, result)
        doc = Document(io.BytesIO(result))
        self.assertEqual([p.text for p in doc.paragraphs],
                         ['Original statement', 'Preserve this separate generation detail'])
        self.assertEqual(doc.tables[0].cell(1, 1).text, 'Sourced agency')

    def test_table_repair_rejects_changed_table_and_missing_evidence(self):
        source = self.fixture()
        plan = self.table_plan(source)
        plan['table_xml_sha256'] = '0' * 64
        with self.assertRaisesRegex(ValueError, 'Table changed'):
            apply_plan(source, plan)
        plan = self.table_plan(source)
        plan['rows'][0]['source_url'] = ''
        with self.assertRaisesRegex(ValueError, 'Missing source provenance'):
            apply_plan(source, plan)


if __name__ == '__main__':
    unittest.main()
