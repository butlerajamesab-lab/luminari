import io
import pathlib
import tempfile
import unittest
import zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree
from audit_docx_corpus import sha, parse_docx, NS
from restore_docx_source_links import restore_links

class SourceLinkRepairTests(unittest.TestCase):
    def fixture(self, linked=False, agency='Agency A'):
        doc=Document();doc.add_paragraph('Original narrative and deadline 30 days.')
        t=doc.add_table(rows=2,cols=3)
        for c,text in zip(t.rows[0].cells,['State','Agency','Website']):c.text=text
        for c,text in zip(t.rows[1].cells,['AL',agency,'Website']):c.text=text
        if linked:
            p=t.cell(1,2).paragraphs[0];h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to('https://official.example.gov/help',RT.HYPERLINK,is_external=True))
            for r in list(p._p):h.append(r)
            p._p.append(h)
        doc.sections[0].footer.paragraphs[0].text='Original footer'
        buf=io.BytesIO();doc.save(buf);return buf.getvalue()

    def planned(self,folder,source,donor):
        dp=pathlib.Path(folder)/'donor.docx';dp.write_bytes(donor)
        loc='/w:document/w:body/w:tbl/w:tr[2]/w:tc[3]'
        return [{'target_sha256':sha(source),'target_location':'word/document.xml:'+loc,'url':'https://official.example.gov/help','donors':[{'source_path':str(dp),'source_sha256':sha(donor),'cell_location':loc}]}]

    def test_restoration_preserves_every_text_and_non_document_part(self):
        src=self.fixture();donor=self.fixture(True)
        with tempfile.TemporaryDirectory() as folder:
            out=restore_links(src,self.planned(folder,src,donor))
        self.assertEqual(parse_docx(src,sha(src))['text_sha256'],parse_docx(out,sha(out))['text_sha256'])
        with zipfile.ZipFile(io.BytesIO(src)) as a,zipfile.ZipFile(io.BytesIO(out)) as b:
            for name in a.namelist():
                if name not in ('word/document.xml','word/_rels/document.xml.rels'):self.assertEqual(a.read(name),b.read(name))
            root=etree.fromstring(b.read('word/document.xml'));self.assertEqual(len(root.xpath('//w:hyperlink',namespaces=NS)),1)

    def test_rejects_similar_but_different_agency(self):
        src=self.fixture()
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(ValueError,'Donor row differs'):restore_links(src,self.planned(folder,src,self.fixture(True,'Agency B')))

    def test_rejects_modified_target(self):
        src=self.fixture()
        with tempfile.TemporaryDirectory() as folder:
            plans=self.planned(folder,src,self.fixture(True))
            with self.assertRaisesRegex(ValueError,'Source bytes'):restore_links(self.fixture(agency='Another generation'),plans)

    def test_rejects_unproven_url(self):
        src=self.fixture()
        with tempfile.TemporaryDirectory() as folder:
            plans=self.planned(folder,src,self.fixture(True));plans[0]['url']='https://unproven.example.org'
            with self.assertRaisesRegex(ValueError,'planned URL'):restore_links(src,plans)

if __name__=='__main__':unittest.main()
